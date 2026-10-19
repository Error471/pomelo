import csv
import html
import io
import re
import shutil
import tempfile
import threading
from pathlib import Path

from ..registry import register
from ..utils import all_ext, ext, run_cmd, soffice_exe, stem, unique_path

WORD_IN = {"doc", "docx", "odt", "rtf"}
TEXT_IN = {"txt", "md", "markdown"}
HTML_IN = {"html", "htm"}
SHEET_IN = {"xlsx", "xlsm", "xls", "ods", "csv", "tsv"}
SLIDE_IN = {"ppt", "pptx", "odp", "pps", "ppsx"}
SUB_IN = {"srt", "vtt"}

_LO_LOCK = threading.Lock()


def _fitz():
    try:
        import pymupdf
        return pymupdf
    except ImportError:
        import fitz
        return fitz


def read_text(path) -> str:
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def not_all(files, same):
    return not all(ext(f) in same for f in files)


# --- LibreOffice ---

def lo_convert(src, target, ctx, filt=None, infilter=None, dest=None):
    exe = soffice_exe()
    if not exe:
        raise RuntimeError("Necesitas LibreOffice instalado")
    profile = Path(tempfile.gettempdir()) / "pomelo-lo-profile"
    with _LO_LOCK, tempfile.TemporaryDirectory() as tmp:
        cmd = [exe, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore",
               "--convert-to", f"{target}:{filt}" if filt else target, "--outdir", tmp]
        if infilter:
            cmd.append(f"--infilter={infilter}")
        cmd.append(str(src))
        run_cmd(cmd)
        produced = [p for p in Path(tmp).iterdir() if p.is_file()]
        if not produced:
            raise RuntimeError("LibreOffice no generó ningún archivo")
        out = Path(dest) if dest else ctx.out(src, target)
        shutil.move(str(produced[0]), out)
    return out


# --- Documentos de texto (Word, ODT, RTF) ---

WORD_TARGETS = [("pdf", "PDF", None), ("docx", "DOCX", "MS Word 2007 XML"), ("odt", "ODT", None),
                ("rtf", "RTF", None), ("txt", "TXT", "Text (encoded):UTF8"), ("html", "HTML", None)]

for _t, _label, _filt in WORD_TARGETS:
    register(id=f"doc.{_t}", label=_label, mode="convert", group="doc", requires=("libreoffice",),
             accepts=lambda f, t=_t: all_ext(f, WORD_IN) and not_all(f, {t}),
             func=lambda src, ctx, t=_t, fl=_filt: lo_convert(src, t, ctx, fl))


def docx_to_md(src, ctx):
    import mammoth
    with open(src, "rb") as fh:
        result = mammoth.convert_to_markdown(fh)
    out = ctx.out(src, "md")
    out.write_text(result.value, encoding="utf-8")
    return out


register(id="doc.md", label="Markdown", mode="convert", group="doc", requires=("py:mammoth",),
         accepts=lambda f: all_ext(f, {"docx"}), func=docx_to_md)


# --- Texto plano, Markdown y HTML ---

HTML_TEMPLATE = """<!doctype html>
<html lang="es"><head><meta charset="utf-8"><title>{title}</title>
<style>body{{font-family:Segoe UI,system-ui,sans-serif;max-width:780px;margin:40px auto;padding:0 16px;line-height:1.6}}
pre,code{{background:#f4f4f4;border-radius:4px}}pre{{padding:12px;overflow:auto}}table{{border-collapse:collapse}}
td,th{{border:1px solid #ccc;padding:4px 8px}}</style></head><body>
{body}
</body></html>
"""


def text_to_html(src, ctx):
    text = read_text(src)
    if ext(src) in ("md", "markdown"):
        import markdown
        body = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])
    else:
        body = f"<pre>{html.escape(text)}</pre>"
    out = ctx.out(src, "html")
    out.write_text(HTML_TEMPLATE.format(title=html.escape(stem(src)), body=body), encoding="utf-8")
    return out


def _qt_document(src):
    from PySide6.QtGui import QFont, QTextDocument
    doc = QTextDocument()
    doc.setDefaultFont(QFont("Segoe UI", 11))
    text = read_text(src)
    e = ext(src)
    if e in ("md", "markdown"):
        doc.setMarkdown(text)
    elif e in HTML_IN:
        doc.setHtml(text)
    else:
        doc.setPlainText(text)
    return doc


def text_to_pdf(src, ctx):
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter
    doc = _qt_document(src)
    out = ctx.out(src, "pdf")
    writer = QPdfWriter(str(out))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageMargins(QMarginsF(18, 18, 18, 18), QPageLayout.Unit.Millimeter)
    writer.setTitle(stem(src))
    doc.print_(writer)
    return out


def html_to_txt(src, ctx):
    out = ctx.out(src, "txt")
    out.write_text(_qt_document(src).toPlainText(), encoding="utf-8")
    return out


_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _add_runs(par, text):
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            par.add_run(part[1:-1]).font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def text_to_docx(src, ctx):
    from docx import Document
    doc = Document()
    is_md = ext(src) in ("md", "markdown")
    for line in read_text(src).splitlines():
        if not line.strip():
            continue
        if is_md:
            if m := re.match(r"^(#{1,6})\s+(.*)", line):
                doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 9))
                continue
            if m := re.match(r"^\s*[-*+]\s+(.*)", line):
                _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
                continue
            if m := re.match(r"^\s*\d+[.)]\s+(.*)", line):
                _add_runs(doc.add_paragraph(style="List Number"), m.group(1))
                continue
            _add_runs(doc.add_paragraph(), line)
        else:
            doc.add_paragraph(line)
    out = ctx.out(src, "docx")
    doc.save(out)
    return out


register(id="text.pdf", label="PDF", mode="convert", group="pdf", main_thread=True,
         accepts=lambda f: all_ext(f, TEXT_IN | HTML_IN), func=text_to_pdf)
register(id="text.html", label="HTML", mode="convert", group="doc", requires=("py:markdown",),
         accepts=lambda f: all_ext(f, TEXT_IN), func=text_to_html)
register(id="text.docx", label="DOCX", mode="convert", group="doc", requires=("py:docx",),
         accepts=lambda f: all_ext(f, TEXT_IN), func=text_to_docx)
register(id="html.txt", label="TXT", mode="convert", group="doc", main_thread=True,
         accepts=lambda f: all_ext(f, HTML_IN), func=html_to_txt)
register(id="html.docx", label="DOCX", mode="convert", group="doc", requires=("libreoffice",),
         accepts=lambda f: all_ext(f, HTML_IN),
         func=lambda src, ctx: lo_convert(src, "docx", ctx, "MS Word 2007 XML", "HTML (StarWriter)"))


# --- PDF ---

def pdf_to_images(pdf, ctx, fmt, name_src=None):
    fitz = _fitz()
    name_src = Path(name_src or pdf)
    dpi = int(ctx.config.get("pdf_dpi", 150))
    with fitz.open(pdf) as doc:
        if doc.page_count == 1:
            out = ctx.out(name_src, fmt)
            doc[0].get_pixmap(dpi=dpi).save(out)
            return out
        folder = unique_path(ctx.folder_for(name_src) / f"{stem(name_src)} (páginas)", is_dir=True)
        folder.mkdir(parents=True)
        width = len(str(doc.page_count))
        for i, page in enumerate(doc, 1):
            page.get_pixmap(dpi=dpi).save(folder / f"{stem(name_src)} - {i:0{width}d}.{fmt}")
        return folder


def pdf_to_txt(src, ctx):
    fitz = _fitz()
    with fitz.open(src) as doc:
        text = "\n\n".join(page.get_text() for page in doc)
    out = ctx.out(src, "txt")
    out.write_text(text, encoding="utf-8")
    return out


def pdf_to_docx(src, ctx):
    from pdf2docx import Converter
    out = ctx.out(src, "docx")
    cv = Converter(str(src))
    try:
        cv.convert(str(out))
    finally:
        cv.close()
    return out


def pdf_to_pptx(src, ctx):
    fitz = _fitz()
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    with fitz.open(src) as doc:
        first = doc[0].rect
        prs.slide_width, prs.slide_height = Emu(int(first.width * 12700)), Emu(int(first.height * 12700))
        blank = prs.slide_layouts[6]
        for page in doc:
            stream = io.BytesIO(page.get_pixmap(dpi=150).tobytes("png"))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(stream, 0, 0, prs.slide_width, prs.slide_height)
    out = ctx.out(src, "pptx")
    prs.save(out)
    return out


_pdf = lambda f: all_ext(f, {"pdf"})  # noqa: E731
register(id="pdf.png", label="PNG", mode="convert", group="image", requires=("py:fitz",),
         accepts=_pdf, func=lambda s, c: pdf_to_images(s, c, "png"))
register(id="pdf.jpg", label="JPG", mode="convert", group="image", requires=("py:fitz",),
         accepts=_pdf, func=lambda s, c: pdf_to_images(s, c, "jpg"))
register(id="pdf.docx", label="DOCX", mode="convert", group="doc", requires=("py:pdf2docx",),
         accepts=_pdf, func=pdf_to_docx)
register(id="pdf.pptx", label="PPTX", mode="convert", group="slides", requires=("py:fitz", "py:pptx"),
         accepts=_pdf, func=pdf_to_pptx)
register(id="pdf.txt", label="TXT", mode="convert", group="doc", requires=("py:fitz",),
         accepts=_pdf, func=pdf_to_txt)


def merge_pdfs(files, ctx):
    fitz = _fitz()
    files = sorted(files, key=lambda f: f.name.lower())
    out = unique_path(ctx.folder_for(files[0]) / f"{stem(files[0])} (unido).pdf")
    with fitz.open() as merged:
        for f in files:
            with fitz.open(f) as part:
                merged.insert_pdf(part)
        merged.save(out, garbage=3, deflate=True)
    return out


def split_pdf(src, ctx):
    fitz = _fitz()
    with fitz.open(src) as doc:
        folder = unique_path(ctx.folder_for(src) / f"{stem(src)} (dividido)", is_dir=True)
        folder.mkdir(parents=True)
        width = len(str(doc.page_count))
        for i in range(doc.page_count):
            with fitz.open() as single:
                single.insert_pdf(doc, from_page=i, to_page=i)
                single.save(folder / f"{stem(src)} - {i + 1:0{width}d}.pdf")
    return folder


def compress_pdf(src, ctx):
    fitz = _fitz()
    out = ctx.out(src, "pdf", " (comprimido)")
    with fitz.open(src) as doc:
        if hasattr(doc, "rewrite_images"):
            try:
                doc.rewrite_images(dpi_threshold=200, dpi_target=150, quality=70)
            except Exception:
                pass
        doc.save(out, garbage=4, deflate=True, clean=True)
    return out


def rotate_pdf(src, ctx):
    fitz = _fitz()
    out = ctx.out(src, "pdf", " (rotado)")
    with fitz.open(src) as doc:
        for page in doc:
            page.set_rotation((page.rotation + 90) % 360)
        doc.save(out, garbage=3, deflate=True)
    return out


register(id="pdf.merge", label="Unir PDFs", mode="tools", group="pdf", requires=("py:fitz",), per_file=False,
         accepts=lambda f: _pdf(f) and len(f) > 1, func=merge_pdfs)
register(id="pdf.split", label="Dividir páginas", mode="tools", group="pdf", requires=("py:fitz",),
         accepts=_pdf, func=split_pdf)
register(id="pdf.compress", label="Comprimir PDF", mode="tools", group="pdf", requires=("py:fitz",),
         accepts=_pdf, func=compress_pdf)
register(id="pdf.rotate", label="Rotar 90°", mode="tools", group="pdf", requires=("py:fitz",),
         accepts=_pdf, func=rotate_pdf)


# --- Hojas de cálculo ---

def _csv_rows(path, delimiter=None):
    text = read_text(path)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    return [_typed(row, delimiter) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _typed(row, delimiter):
    out = []
    for v in row:
        s = v.strip()
        if re.fullmatch(r"-?\d{1,15}", s) and not (len(s.lstrip("-")) > 1 and s.lstrip("-").startswith("0")):
            out.append(int(s))
        elif re.fullmatch(r"-?\d+\.\d+", s):
            out.append(float(s))
        elif delimiter != "," and re.fullmatch(r"-?\d+,\d+", s):
            out.append(float(s.replace(",", ".")))
        else:
            out.append(v)
    return out


def read_sheets(src, ctx):
    import openpyxl
    e = ext(src)
    if e in ("csv", "tsv"):
        return [(stem(src)[:31] or "Hoja1", _csv_rows(src, "\t" if e == "tsv" else None))]
    if e in ("xls", "ods"):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_x = lo_convert(src, "xlsx", ctx, "Calc MS Excel 2007 XML", dest=Path(tmp) / "t.xlsx")
            return _read_xlsx(openpyxl, tmp_x)
    return _read_xlsx(openpyxl, src)


def _read_xlsx(openpyxl, path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def write_xlsx(sheets, out):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", str(title))[:31] or "Hoja")
        for row in rows:
            ws.append(row)
    wb.save(out)
    return out


def sheet_to_xlsx(src, ctx):
    return write_xlsx(read_sheets(src, ctx), ctx.out(src, "xlsx"))


def sheet_to_delimited(src, ctx, target):
    sheets = read_sheets(src, ctx)
    delim = "\t" if target == "tsv" else ctx.config.get("csv_separator", ";")
    outs = []
    for title, rows in sheets:
        suffix = f" - {title}" if len(sheets) > 1 else ""
        out = ctx.out(src, target, suffix)
        with open(out, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh, delimiter=delim)
            for row in rows:
                w.writerow(["" if v is None else v for v in row])
        outs.append(out)
    return outs


def sheet_via_lo(src, ctx, target):
    if ext(src) in ("csv", "tsv"):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_x = write_xlsx(read_sheets(src, ctx), Path(tmp) / f"{stem(src)}.xlsx")
            return lo_convert(tmp_x, target, ctx, dest=ctx.out(src, target))
    return lo_convert(src, target, ctx)


_sheet = lambda f, same: all_ext(f, SHEET_IN) and not_all(f, same)  # noqa: E731
register(id="sheet.xlsx", label="XLSX", mode="convert", group="sheet", requires=("py:openpyxl",),
         accepts=lambda f: _sheet(f, {"xlsx"}), func=sheet_to_xlsx)
register(id="sheet.csv", label="CSV", mode="convert", group="sheet", requires=("py:openpyxl",),
         accepts=lambda f: _sheet(f, {"csv"}), func=lambda s, c: sheet_to_delimited(s, c, "csv"))
register(id="sheet.tsv", label="TSV", mode="convert", group="sheet", requires=("py:openpyxl",),
         accepts=lambda f: _sheet(f, {"tsv"}), func=lambda s, c: sheet_to_delimited(s, c, "tsv"))
register(id="sheet.ods", label="ODS", mode="convert", group="sheet", requires=("libreoffice",),
         accepts=lambda f: _sheet(f, {"ods"}), func=lambda s, c: sheet_via_lo(s, c, "ods"))
register(id="sheet.pdf", label="PDF", mode="convert", group="pdf", requires=("libreoffice",),
         accepts=lambda f: all_ext(f, SHEET_IN), func=lambda s, c: sheet_via_lo(s, c, "pdf"))


# --- Presentaciones ---

def slides_to_images(src, ctx, fmt):
    with tempfile.TemporaryDirectory() as tmp:
        pdf = lo_convert(src, "pdf", ctx, dest=Path(tmp) / "slides.pdf")
        return pdf_to_images(pdf, ctx, fmt, name_src=src)


for _t, _label in (("pdf", "PDF"), ("pptx", "PPTX"), ("odp", "ODP")):
    register(id=f"slides.{_t}", label=_label, mode="convert", group="pdf" if _t == "pdf" else "slides",
             requires=("libreoffice",), accepts=lambda f, t=_t: all_ext(f, SLIDE_IN) and not_all(f, {t}),
             func=lambda s, c, t=_t: lo_convert(s, t, c))
for _t in ("png", "jpg"):
    register(id=f"slides.{_t}", label=_t.upper(), mode="convert", group="image",
             requires=("libreoffice", "py:fitz"), accepts=lambda f: all_ext(f, SLIDE_IN),
             func=lambda s, c, t=_t: slides_to_images(s, c, t))


# --- Subtítulos ---

_TS = re.compile(r"(?:(\d+):)?(\d{1,2}):(\d{2})[.,](\d{3})")


def parse_subs(text):
    cues = []
    for block in re.split(r"\n\s*\n", text.replace("\r", "")):
        lines = [ln for ln in block.split("\n") if ln.strip()]
        idx = next((i for i, ln in enumerate(lines) if "-->" in ln), None)
        if idx is None:
            continue
        times = _TS.findall(lines[idx])
        if len(times) < 2:
            continue
        start, end = (((int(h or 0) * 60 + int(m)) * 60 + int(s)) * 1000 + int(ms) for h, m, s, ms in times[:2])
        cues.append((start, end, lines[idx + 1:]))
    return cues


def _fmt_ts(ms, sep):
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return f"{h:02}:{m:02}:{s:02}{sep}{ms:03}"


def convert_subs(src, ctx, target):
    cues = parse_subs(read_text(src))
    if target == "srt":
        body = "\n".join(f"{i}\n{_fmt_ts(a, ',')} --> {_fmt_ts(b, ',')}\n" + "\n".join(t) + "\n"
                         for i, (a, b, t) in enumerate(cues, 1))
    elif target == "vtt":
        body = "WEBVTT\n\n" + "\n".join(f"{_fmt_ts(a, '.')} --> {_fmt_ts(b, '.')}\n" + "\n".join(t) + "\n"
                                        for a, b, t in cues)
    else:
        body = "\n".join(re.sub(r"<[^>]+>", "", " ".join(t)) for _, _, t in cues) + "\n"
    out = ctx.out(src, target)
    out.write_text(body, encoding="utf-8")
    return out


for _t, _label in (("srt", "SRT"), ("vtt", "VTT"), ("txt", "TXT")):
    register(id=f"subs.{_t}", label=_label, mode="convert", group="doc",
             accepts=lambda f, t=_t: all_ext(f, SUB_IN) and not_all(f, {t}),
             func=lambda s, c, t=_t: convert_subs(s, c, t))
